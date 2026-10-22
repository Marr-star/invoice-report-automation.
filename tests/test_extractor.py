from decimal import Decimal
from pathlib import Path
from invoice_automation.extractor import parse_invoice
from invoice_automation.report import process_folder

BASE = "Vendor: Acme Tools\nInvoice No: A-100\nInvoice Date: 2026-10-21\nCurrency: USD\nSubtotal: 100.00\nTax: 14.00\nGrand Total: 114.00\n"


def test_valid_invoice():
    record, errors = parse_invoice(BASE)
    assert not errors
    assert record["total"] == Decimal("114.00")
    assert record["invoice_date"] == "2026-10-21"


def test_missing_field():
    _, errors = parse_invoice(BASE.replace("Vendor: Acme Tools\n", ""))
    assert "missing_or_ambiguous_vendor" in errors


def test_invalid_date():
    _, errors = parse_invoice(BASE.replace("2026-10-21", "2026-99-99"))
    assert "invalid_invoice_date" in errors


def test_inconsistent_total():
    _, errors = parse_invoice(BASE.replace("Grand Total: 114.00", "Grand Total: 120.00"))
    assert "total_does_not_match_subtotal_plus_tax" in errors


def test_empty_scanned_like():
    _, errors = parse_invoice("   ")
    assert "no_extractable_text_scanned_pdf_possible" in errors


def test_duplicate_invoice_and_multiple_files(tmp_path):
    from reportlab.pdfgen import canvas
    pdfs = tmp_path / "pdfs"
    pdfs.mkdir()
    for name in ("a.pdf", "b.pdf"):
        c = canvas.Canvas(str(pdfs / name))
        y = 780
        for line in BASE.splitlines():
            c.drawString(72, y, line)
            y -= 22
        c.save()
    result = process_folder(pdfs, tmp_path / "out")
    assert result["processed"] == 2
    assert result["accepted"] == 1
    assert result["review"] == 1
    assert "duplicate_vendor_invoice_id" in (tmp_path / "out" / "processing_errors.csv").read_text()


def test_no_pdfs(tmp_path):
    import pytest
    with pytest.raises(ValueError, match="No PDFs"):
        process_folder(tmp_path, tmp_path / "output")


import pytest


@pytest.mark.parametrize('value', ['114.001', '114oops', '1,14.00', 'NaN', 'Infinity', '114.00 USD'])
def test_malformed_money_is_not_partially_accepted(value):
    _, errors = parse_invoice(BASE.replace('Grand Total: 114.00', f'Grand Total: {value}'))
    assert 'invalid_total' in errors


@pytest.mark.parametrize('value', ['-114.00', '-$114.00'])
def test_negative_total(value):
    _, errors = parse_invoice(BASE.replace('Grand Total: 114.00', f'Grand Total: {value}'))
    assert 'negative_total' in errors


def test_ambiguous_total():
    _, errors = parse_invoice(BASE + 'Total: 999.00\n')
    assert 'missing_or_ambiguous_total' in errors


def test_multiple_invoice_ids_in_one_pdf_require_review():
    _, errors = parse_invoice(BASE + 'Invoice No: A-101\n')
    assert 'missing_or_ambiguous_invoice_id' in errors


def test_identifier_must_be_valid():
    _, errors = parse_invoice(BASE.replace('A-100', 'A 100?'))
    assert 'invalid_invoice_id' in errors


def test_currency_symbol_mismatch():
    _, errors = parse_invoice(BASE.replace('114.00', '€114.00'))
    assert 'currency_symbol_mismatch_total' in errors


def test_unsupported_currency():
    _, errors = parse_invoice(BASE.replace('USD', 'XYZ'))
    assert 'unsupported_currency' in errors


def test_day_first_date_and_vendor_whitespace():
    record, errors = parse_invoice(BASE.replace('2026-10-21', '21/10/2026').replace('Acme Tools', 'Acme   Tools'))
    assert not errors
    assert record['vendor'] == 'Acme Tools'
    assert record['invoice_date'] == '2026-10-21'


def test_batch_outputs_and_real_excel_numeric_cells(tmp_path):
    from openpyxl import load_workbook
    import pandas as pd
    result = process_folder(Path(__file__).parents[1] / 'samples/invoices', tmp_path)
    assert (result['processed'], result['accepted'], result['review']) == (13, 9, 4)
    rows = pd.read_csv(tmp_path / 'invoices.csv')
    assert len(rows) == 9
    assert result['totals_by_currency'] == {'EGP': '855.00', 'EUR': '684.00', 'USD': '1026.00'}
    book = load_workbook(tmp_path / 'invoice_report.xlsx')
    assert book.sheetnames == ['Invoices', 'Needs Review', 'Currency Summary']
    assert book['Invoices']['H2'].data_type == 'n'
    assert book['Invoices'].freeze_panes == 'A2'
    assert (tmp_path / 'processing_summary.json').exists()


def test_corrupt_pdf_is_reported(tmp_path):
    folder = tmp_path / 'in'
    folder.mkdir()
    (folder / 'bad.PDF').write_bytes(b'not a pdf')
    result = process_folder(folder, tmp_path / 'out')
    assert result['accepted'] == 0
    assert result['review'] == 1
    assert 'pdf_read_error' in (tmp_path / 'out/processing_errors.csv').read_text()


def test_formula_injection_is_escaped(tmp_path):
    from reportlab.pdfgen import canvas
    from openpyxl import load_workbook
    folder = tmp_path / 'in'
    folder.mkdir()
    c = canvas.Canvas(str(folder / 'sample.pdf'))
    for i, line in enumerate(BASE.replace('Acme Tools', '=HYPERLINK("bad")').splitlines()):
        c.drawString(72, 780 - i * 25, line)
    c.save()
    process_folder(folder, tmp_path / 'out')
    book = load_workbook(tmp_path / 'out/invoice_report.xlsx')
    assert book['Invoices']['B2'].data_type == 's'
    assert book['Invoices']['B2'].value.startswith("'=")


def test_all_invalid_produces_headers_and_empty_summary(tmp_path):
    folder = tmp_path / 'in'
    folder.mkdir()
    import fitz
    doc = fitz.open()
    doc.new_page()
    doc.save(folder / 'blank.pdf')
    doc.close()
    result = process_folder(folder, tmp_path / 'out')
    assert result == {'processed': 1, 'accepted': 0, 'review': 1, 'totals_by_currency': {}}
    assert (tmp_path / 'out/invoices.csv').read_text().startswith('source_file,vendor,invoice_id')
